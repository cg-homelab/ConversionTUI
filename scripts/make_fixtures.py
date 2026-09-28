"""Regenerate the tiny test fixtures in tests/fixtures.

Run with `uv run python scripts/make_fixtures.py`. The generated files are
committed, so this only needs running when a fixture changes — but keeping the
generator means nobody has to reverse-engineer a binary blob later.

Deliberately hand-rolled: markitdown depends on mammoth (not python-docx) and
brings no PDF writer, so authoring these with a library would mean adding dev
dependencies purely to produce a few kilobytes.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

FIXTURES = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
MARKER = "Conversion fixture heading"
BODY = "The quick brown fox jumps over the lazy dog."


def write_docx(path: Path) -> None:
    """A minimal OOXML word document: heading plus one paragraph."""
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    document = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{MARKER}</w:t></w:r></w:p>
<w:p><w:r><w:t>{BODY}</w:t></w:r></w:p>
</w:body>
</w:document>"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)


def write_pdf(path: Path) -> None:
    """A minimal single-page PDF with one line of extractable text."""
    text = f"{MARKER} {BODY}"
    stream = f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def write_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = MARKER
    slide.placeholders[1].text = BODY
    prs.save(path)


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["heading", "body"])
    ws.append([MARKER, BODY])
    wb.save(path)


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    write_docx(FIXTURES / "sample.docx")
    write_pdf(FIXTURES / "sample.pdf")
    write_pptx(FIXTURES / "sample.pptx")
    write_xlsx(FIXTURES / "sample.xlsx")
    (FIXTURES / "sample.csv").write_text(f"heading,body\n{MARKER},{BODY}\n", encoding="utf-8")
    (FIXTURES / "sample.html").write_text(
        f"<html><body><h1>{MARKER}</h1><p>{BODY}</p></body></html>\n", encoding="utf-8"
    )
    # Valid extension, garbage bytes: asserts a clean failure, not a traceback.
    (FIXTURES / "corrupt.pdf").write_bytes(b"%PDF-1.4\nthis is not a pdf at all\n")

    for f in sorted(FIXTURES.iterdir()):
        if f.is_file():
            print(f"{f.name:16} {f.stat().st_size:>7} bytes")


if __name__ == "__main__":
    main()
